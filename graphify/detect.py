# file discovery, type classification, and corpus health checks
from __future__ import annotations
import fnmatch
import json
import os
import re
from dataclasses import dataclass
from enum import Enum
from pathlib import Path


class FileType(str, Enum):
    CODE = "code"
    DOCUMENT = "document"
    PAPER = "paper"
    IMAGE = "image"


_MANIFEST_PATH = "graphify-out/manifest.json"

CODE_EXTENSIONS = {'.py', '.ts', '.js', '.tsx', '.go', '.rs', '.java', '.cpp', '.cc', '.cxx', '.c', '.h', '.hpp', '.rb', '.swift', '.kt', '.kts', '.cs', '.scala', '.php', '.lua', '.toc', '.zig', '.ps1', '.ex', '.exs', '.m', '.mm'}
DOC_EXTENSIONS = {'.md', '.txt', '.rst'}
PAPER_EXTENSIONS = {'.pdf'}
IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.svg'}
OFFICE_EXTENSIONS = {'.docx', '.xlsx'}

CORPUS_WARN_THRESHOLD = 50_000    # words - below this, warn "you may not need a graph"
CORPUS_UPPER_THRESHOLD = 500_000  # words - above this, warn about token cost
FILE_COUNT_UPPER = 200             # files - above this, warn about token cost

# Files that may contain secrets - skip silently
_SENSITIVE_PATTERNS = [
    re.compile(r'(^|[\\/])\.(env|envrc)(\.|$)', re.IGNORECASE),
    re.compile(r'\.(pem|key|p12|pfx|cert|crt|der|p8)$', re.IGNORECASE),
    re.compile(r'(credential|secret|passwd|password|token|private_key)', re.IGNORECASE),
    re.compile(r'(id_rsa|id_dsa|id_ecdsa|id_ed25519)(\.pub)?$'),
    re.compile(r'(\.netrc|\.pgpass|\.htpasswd)$', re.IGNORECASE),
    re.compile(r'(aws_credentials|gcloud_credentials|service.account)', re.IGNORECASE),
]

# Signals that a .md/.txt file is actually a converted academic paper
_PAPER_SIGNALS = [
    re.compile(r'\barxiv\b', re.IGNORECASE),
    re.compile(r'\bdoi\s*:', re.IGNORECASE),
    re.compile(r'\babstract\b', re.IGNORECASE),
    re.compile(r'\bproceedings\b', re.IGNORECASE),
    re.compile(r'\bjournal\b', re.IGNORECASE),
    re.compile(r'\bpreprint\b', re.IGNORECASE),
    re.compile(r'\\cite\{'),          # LaTeX citation
    re.compile(r'\[\d+\]'),           # Numbered citation [1], [23] (inline)
    re.compile(r'\[\n\d+\n\]'),       # Numbered citation spread across lines (markdown conversion)
    re.compile(r'eq\.\s*\d+|equation\s+\d+', re.IGNORECASE),
    re.compile(r'\d{4}\.\d{4,5}'),   # arXiv ID like 1706.03762
    re.compile(r'\bwe propose\b', re.IGNORECASE),   # common academic phrasing
    re.compile(r'\bliterature\b', re.IGNORECASE),   # "from the literature"
]
_PAPER_SIGNAL_THRESHOLD = 3  # need at least this many signals to call it a paper


def _is_sensitive(path: Path) -> bool:
    """Return True if this file likely contains secrets and should be skipped."""
    name = path.name
    full = str(path)
    return any(p.search(name) or p.search(full) for p in _SENSITIVE_PATTERNS)


def _looks_like_paper(path: Path) -> bool:
    """Heuristic: does this text file read like an academic paper?"""
    try:
        # Only scan first 3000 chars for speed
        text = path.read_text(errors="ignore")[:3000]
        hits = sum(1 for pattern in _PAPER_SIGNALS if pattern.search(text))
        return hits >= _PAPER_SIGNAL_THRESHOLD
    except Exception:
        return False


def classify_file(path: Path) -> FileType | None:
    ext = path.suffix.lower()
    if ext in CODE_EXTENSIONS:
        return FileType.CODE
    if ext in PAPER_EXTENSIONS:
        return FileType.PAPER
    if ext in IMAGE_EXTENSIONS:
        return FileType.IMAGE
    if ext in DOC_EXTENSIONS:
        # Check if it's a converted paper
        if _looks_like_paper(path):
            return FileType.PAPER
        return FileType.DOCUMENT
    if ext in OFFICE_EXTENSIONS:
        return FileType.DOCUMENT
    return None


def extract_pdf_text(path: Path) -> str:
    """Extract plain text from a PDF file using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n".join(pages)
    except Exception:
        return ""


def docx_to_markdown(path: Path) -> str:
    """Convert a .docx file to markdown text using python-docx."""
    try:
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(str(path))
        lines = []
        for para in doc.paragraphs:
            style = para.style.name if para.style else ""
            text = para.text.strip()
            if not text:
                lines.append("")
                continue
            if style.startswith("Heading 1"):
                lines.append(f"# {text}")
            elif style.startswith("Heading 2"):
                lines.append(f"## {text}")
            elif style.startswith("Heading 3"):
                lines.append(f"### {text}")
            elif style.startswith("List"):
                lines.append(f"- {text}")
            else:
                lines.append(text)
        # Tables
        for table in doc.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            header = "| " + " | ".join(rows[0]) + " |"
            sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
            lines.extend([header, sep])
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")
        return "\n".join(lines)
    except ImportError:
        return ""
    except Exception:
        return ""


def xlsx_to_markdown(path: Path) -> str:
    """Convert an .xlsx file to markdown text using openpyxl."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        sections = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                # Skip entirely empty rows
                if all(cell is None for cell in row):
                    continue
                rows.append([str(cell) if cell is not None else "" for cell in row])
            if not rows:
                continue
            sections.append(f"## Sheet: {sheet_name}")
            if len(rows) >= 1:
                header = "| " + " | ".join(rows[0]) + " |"
                sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
                sections.extend([header, sep])
                for row in rows[1:]:
                    sections.append("| " + " | ".join(row) + " |")
        wb.close()
        return "\n".join(sections)
    except ImportError:
        return ""
    except Exception:
        return ""


def convert_office_file(path: Path, out_dir: Path) -> Path | None:
    """Convert a .docx or .xlsx to a markdown sidecar in out_dir.

    Returns the path of the converted .md file, or None if conversion failed
    or the required library is not installed.
    """
    ext = path.suffix.lower()
    if ext == ".docx":
        text = docx_to_markdown(path)
    elif ext == ".xlsx":
        text = xlsx_to_markdown(path)
    else:
        return None

    if not text.strip():
        return None

    out_dir.mkdir(parents=True, exist_ok=True)
    # Use a stable name derived from the original path to avoid collisions
    import hashlib
    name_hash = hashlib.sha256(str(path.resolve()).encode()).hexdigest()[:8]
    out_path = out_dir / f"{path.stem}_{name_hash}.md"
    out_path.write_text(
        f"<!-- converted from {path.name} -->\n\n{text}",
        encoding="utf-8",
    )
    return out_path


def count_words(path: Path) -> int:
    try:
        ext = path.suffix.lower()
        if ext == ".pdf":
            return len(extract_pdf_text(path).split())
        if ext == ".docx":
            return len(docx_to_markdown(path).split())
        if ext == ".xlsx":
            return len(xlsx_to_markdown(path).split())
        return len(path.read_text(errors="ignore").split())
    except Exception:
        return 0


# Directory names to always skip - venvs, caches, build artifacts, deps
_SKIP_DIRS = {
    "venv", ".venv", "env", ".env",
    "node_modules", "__pycache__", ".git",
    "dist", "build", "target", "out",
    "site-packages", "lib64",
    ".pytest_cache", ".mypy_cache", ".ruff_cache",
    ".tox", ".eggs", "*.egg-info",
}

def _is_noise_dir(part: str) -> bool:
    """Return True if this directory name looks like a venv, cache, or dep dir."""
    if part in _SKIP_DIRS:
        return True
    # Catch *_venv, *_repo/site-packages patterns
    if part.endswith("_venv") or part.endswith("_env"):
        return True
    if part.endswith(".egg-info"):
        return True
    return False


def _load_ignore_file(path: Path) -> list[str]:
    """Read an ignore file (e.g. .graphifyignore, .gitignore) and return patterns.

    Lines starting with # are comments. Blank lines are ignored.
    Patterns follow gitignore semantics: glob matched against the path
    relative to root. A leading slash anchors to root. A trailing slash
    matches directories only (we match both dir and file for simplicity).
    """
    if not path.exists():
        return []
    patterns = []
    for line in path.read_text(errors="ignore").splitlines():
        line = line.strip()
        if line and not line.startswith("#"):
            patterns.append(line)
    return patterns


def _load_graphifyignore(root: Path) -> list[str]:
    """Read .graphifyignore from root and return a list of patterns."""
    return _load_ignore_file(root / ".graphifyignore")


@dataclass(frozen=True)
class GitIgnoreRule:
    base_dir: Path
    pattern: str
    negated: bool


def _load_gitignore(root: Path) -> list[GitIgnoreRule]:
    """Read .gitignore from root and preserve ordered negation rules."""
    return _load_gitignore_file(root / ".gitignore")


def _load_gitignore_file(path: Path) -> list[GitIgnoreRule]:
    """Read one .gitignore file and bind rules to its containing directory."""
    rules: list[GitIgnoreRule] = []
    for line in _load_ignore_file(path):
        negated = line.startswith("!") and len(line) > 1
        pattern = line[1:] if negated else line
        if pattern:
            rules.append(
                GitIgnoreRule(
                    base_dir=path.parent,
                    pattern=pattern,
                    negated=negated,
                )
            )
    return rules


def _matches_ignore_pattern(path: Path, root: Path, pattern: str, *, base_dir: Path | None = None) -> bool:
    """Return True if a path matches one ignore pattern."""
    base_path = base_dir or root
    try:
        rel = str(path.relative_to(base_path))
    except ValueError:
        return False
    rel = rel.replace(os.sep, "/")
    parts = rel.split("/")
    anchored = pattern.startswith("/")
    directory_pattern = pattern.endswith("/")
    normalized = pattern.strip("/")
    if not normalized:
        return False

    if anchored:
        # Leading "/" is anchored to the directory containing the .gitignore file.
        if fnmatch.fnmatch(rel, normalized):
            return True
        if directory_pattern:
            for i in range(len(parts)):
                if fnmatch.fnmatch("/".join(parts[: i + 1]), normalized):
                    return True
        return False

    if "/" in normalized and fnmatch.fnmatch(rel, normalized):
        return True
    if "/" not in normalized and fnmatch.fnmatch(path.name, normalized):
        return True
    for i, part in enumerate(parts):
        if "/" not in normalized and fnmatch.fnmatch(part, normalized):
            return True
        if fnmatch.fnmatch("/".join(parts[:i + 1]), normalized):
            return True
    return False


def _is_ignored(path: Path, root: Path, patterns: list[str]) -> bool:
    """Return True if path matches any .graphifyignore pattern."""
    return any(_matches_ignore_pattern(path, root, pattern) for pattern in patterns)


def _is_gitignored(path: Path, root: Path, rules: list[GitIgnoreRule]) -> bool:
    """Return True if ordered .gitignore rules exclude this path."""
    ignored = False
    for rule in rules:
        if _matches_ignore_pattern(path, root, rule.pattern, base_dir=rule.base_dir):
            ignored = not rule.negated
    return ignored


def detect(root: Path, *, follow_symlinks: bool = False) -> dict:
    files: dict[FileType, list[str]] = {
        FileType.CODE: [],
        FileType.DOCUMENT: [],
        FileType.PAPER: [],
        FileType.IMAGE: [],
    }
    total_words = 0

    skipped_sensitive: list[str] = []
    graphifyignore_patterns = _load_graphifyignore(root)
    root_gitignore_rules = _load_gitignore(root)

    # Always include graphify-out/memory/ - query results filed back into the graph
    memory_dir = root / "graphify-out" / "memory"
    scan_paths = [root]
    if memory_dir.exists():
        scan_paths.append(memory_dir)

    seen: set[Path] = set()
    all_files: list[tuple[Path, list[GitIgnoreRule]]] = []
    rules_by_dir: dict[Path, list[GitIgnoreRule]] = {root: root_gitignore_rules}

    for scan_root in scan_paths:
        in_memory_tree = memory_dir.exists() and str(scan_root).startswith(str(memory_dir))
        for dirpath, dirnames, filenames in os.walk(scan_root, followlinks=follow_symlinks):
            dp = Path(dirpath)
            active_gitignore_rules = rules_by_dir.get(dp, root_gitignore_rules)
            local_gitignore = dp / ".gitignore"
            if local_gitignore.exists() and dp != root:
                active_gitignore_rules = [
                    *active_gitignore_rules,
                    *_load_gitignore_file(local_gitignore),
                ]
                rules_by_dir[dp] = active_gitignore_rules
            if follow_symlinks and os.path.islink(dirpath):
                real = os.path.realpath(dirpath)
                parent_real = os.path.realpath(os.path.dirname(dirpath))
                if parent_real == real or parent_real.startswith(real + os.sep):
                    dirnames.clear()
                    continue
            if not in_memory_tree:
                # Prune noise dirs in-place so os.walk never descends into them
                dirnames[:] = [
                    d for d in dirnames
                    if not d.startswith(".")
                    and not _is_noise_dir(d)
                    and not _is_ignored(dp / d, root, graphifyignore_patterns)
                    and not _is_gitignored(dp / d, root, active_gitignore_rules)
                ]
                for dirname in dirnames:
                    child_dir = dp / dirname
                    rules_by_dir[child_dir] = active_gitignore_rules
            for fname in filenames:
                p = dp / fname
                if p not in seen:
                    seen.add(p)
                    all_files.append((p, active_gitignore_rules))

    converted_dir = root / "graphify-out" / "converted"

    for p, gitignore_rules in all_files:
        # For memory dir files, skip hidden/noise filtering
        in_memory = memory_dir.exists() and str(p).startswith(str(memory_dir))
        if not in_memory:
            # Hidden files are already excluded via dir pruning above,
            # but catch hidden files at the root level
            if p.name.startswith("."):
                continue
            # Skip files inside our own converted/ dir (avoid re-processing sidecars)
            if str(p).startswith(str(converted_dir)):
                continue
        if _is_ignored(p, root, graphifyignore_patterns):
            continue
        if not in_memory and _is_gitignored(p, root, gitignore_rules):
            continue
        if _is_sensitive(p):
            skipped_sensitive.append(str(p))
            continue
        ftype = classify_file(p)
        if ftype:
            # Office files: convert to markdown sidecar so subagents can read them
            if p.suffix.lower() in OFFICE_EXTENSIONS:
                md_path = convert_office_file(p, converted_dir)
                if md_path:
                    files[ftype].append(str(md_path))
                    total_words += count_words(md_path)
                else:
                    # Conversion failed (library not installed) - skip with note
                    skipped_sensitive.append(str(p) + " [office conversion failed - pip install graphifyy[office]]")
                continue
            files[ftype].append(str(p))
            total_words += count_words(p)

    total_files = sum(len(v) for v in files.values())
    needs_graph = total_words >= CORPUS_WARN_THRESHOLD

    # Determine warning - lower bound, upper bound, or sensitive files skipped
    warning: str | None = None
    if not needs_graph:
        warning = (
            f"Corpus is ~{total_words:,} words - fits in a single context window. "
            f"You may not need a graph."
        )
    elif total_words >= CORPUS_UPPER_THRESHOLD or total_files >= FILE_COUNT_UPPER:
        warning = (
            f"Large corpus: {total_files} files · ~{total_words:,} words. "
            f"Semantic extraction will be expensive (many Claude tokens). "
            f"Consider running on a subfolder, or use --no-semantic to run AST-only."
        )

    return {
        "files": {k.value: v for k, v in files.items()},
        "total_files": total_files,
        "total_words": total_words,
        "needs_graph": needs_graph,
        "warning": warning,
        "skipped_sensitive": skipped_sensitive,
        "graphifyignore_patterns": len(graphifyignore_patterns),
        "gitignore_loaded": (root / ".gitignore").exists(),
    }


def load_manifest(manifest_path: str = _MANIFEST_PATH) -> dict[str, float]:
    """Load the file modification time manifest from a previous run."""
    try:
        return json.loads(Path(manifest_path).read_text())
    except Exception:
        return {}


def save_manifest(files: dict[str, list[str]], manifest_path: str = _MANIFEST_PATH) -> None:
    """Save current file mtimes so the next --update run can diff against them."""
    manifest: dict[str, float] = {}
    for file_list in files.values():
        for f in file_list:
            try:
                manifest[f] = Path(f).stat().st_mtime
            except OSError:
                pass  # file deleted between detect() and manifest write - skip it
    Path(manifest_path).parent.mkdir(parents=True, exist_ok=True)
    Path(manifest_path).write_text(json.dumps(manifest, indent=2))


def detect_incremental(root: Path, manifest_path: str = _MANIFEST_PATH) -> dict:
    """Like detect(), but returns only new or modified files since the last run.

    Compares current file mtimes against the stored manifest.
    Use for --update mode: re-extract only what changed, merge into existing graph.
    """
    full = detect(root)
    manifest = load_manifest(manifest_path)

    if not manifest:
        # No previous run - treat everything as new
        full["incremental"] = True
        full["new_files"] = full["files"]
        full["unchanged_files"] = {k: [] for k in full["files"]}
        full["new_total"] = full["total_files"]
        return full

    new_files: dict[str, list[str]] = {k: [] for k in full["files"]}
    unchanged_files: dict[str, list[str]] = {k: [] for k in full["files"]}

    for ftype, file_list in full["files"].items():
        for f in file_list:
            stored_mtime = manifest.get(f)
            try:
                current_mtime = Path(f).stat().st_mtime
            except Exception:
                current_mtime = 0
            if stored_mtime is None or current_mtime > stored_mtime:
                new_files[ftype].append(f)
            else:
                unchanged_files[ftype].append(f)

    # Files in manifest that no longer exist - their cached nodes are now ghost nodes
    current_files = {f for flist in full["files"].values() for f in flist}
    deleted_files = [f for f in manifest if f not in current_files]

    new_total = sum(len(v) for v in new_files.values())
    full["incremental"] = True
    full["new_files"] = new_files
    full["unchanged_files"] = unchanged_files
    full["new_total"] = new_total
    full["deleted_files"] = deleted_files
    return full
